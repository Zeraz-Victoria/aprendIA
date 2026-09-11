import os
import json
import openpyxl

EXCEL_2 = '/Users/adrianvira/Library/Mobile Documents/com~apple~CloudDocs/Ciclo escolar 2025-2026/Fase 6/Fase 6_2\xba_Contenidos y PDA.xlsx'
OUTPUT_PATH = '/Users/adrianvira/.gemini/antigravity-ide/scratch/aprendIA/data/contenidos_pda.json'

SHEET_MAP = {
    ' 2\xb0 ESP':    {'materia': '2\xb0 ESP',      'campo': 'LENGUAJES'},
    '2\xb0ING':      {'materia': '2\xb0 ING',      'campo': 'LENGUAJES'},
    '2\xb0ARTES':    {'materia': '2\xb0 ART',      'campo': 'LENGUAJES'},
    '2\xb0MAT':      {'materia': '2\xb0 MAT',      'campo': 'SABERES Y PENSAMIENTO CIENT\xcdFICO'},
    '2\xb0FIS':      {'materia': '2\xb0 FIS',      'campo': 'SABERES Y PENSAMIENTO CIENT\xcdFICO'},
    '2\xb0HIST':     {'materia': '2\xb0 HIST',     'campo': '\xc9TICA, NATURALEZA Y SOCIEDADES'},
    '2\xb0F.C.E':    {'materia': '2\xb0 F.C.E',   'campo': '\xc9TICA, NATURALEZA Y SOCIEDADES'},
    '2\xb0TEC':      {'materia': '2\xb0 TEC',      'campo': 'DE LO HUMANO Y LO COMUNITARIO'},
    '2\xb0TUTO':     {'materia': '2\xb0 TUTO',     'campo': 'DE LO HUMANO Y LO COMUNITARIO'},
    '2\xb0EDU.FIS.': {'materia': '2\xb0 EDU.FIS', 'campo': 'DE LO HUMANO Y LO COMUNITARIO'},
}

SKIP_KEYWORDS = [
    'CAMPO FORMATIVO', 'CONTENIDOS DE LA FASE', 'PROCESOS DE DESARROLLO',
    'ZONAESCOLAR', 'HTTP', 'GRADO', 'PDA', 'ANDR\xc9S'
]

def should_skip(text):
    if not text:
        return True
    t = str(text).strip()
    if len(t) < 5:
        return True
    tu = t.upper()
    for kw in SKIP_KEYWORDS:
        if kw in tu:
            return True
    return False

def extract_records(ws, meta):
    records = []
    curr_contenido = None
    for row in ws.iter_rows(values_only=True):
        col0 = str(row[0]).strip() if row[0] else ''
        pda_val = ''
        if len(row) > 2 and row[2]:
            pda_val = str(row[2]).strip()
        elif len(row) > 1 and row[1]:
            pda_val = str(row[1]).strip()

        if should_skip(col0) and should_skip(pda_val):
            continue

        if col0 and not should_skip(col0):
            curr_contenido = col0

        if pda_val and not should_skip(pda_val) and curr_contenido:
            for pda in [p.strip() for p in pda_val.split('\n') if p.strip() and len(p.strip()) > 10]:
                records.append({
                    'grado': '2\xb0 Secundaria',
                    'fase': 'Fase 6',
                    'materia': meta['materia'],
                    'campo_formativo': meta['campo'],
                    'contenido': curr_contenido,
                    'pda': pda
                })
    return records

def main():
    existing_1 = []
    if os.path.exists(OUTPUT_PATH):
        with open(OUTPUT_PATH, 'r', encoding='utf-8') as f:
            existing = json.load(f)
        existing_1 = [r for r in existing if r.get('grado') == '1\xb0 Secundaria']
        print(f'Conservando {len(existing_1)} registros de 1\xb0 Secundaria')

    wb = openpyxl.load_workbook(EXCEL_2, data_only=True)
    new_records = []

    for sn in wb.sheetnames:
        meta = SHEET_MAP.get(sn)
        if not meta:
            print(f'Hoja no mapeada: {sn}')
            continue
        records = extract_records(wb[sn], meta)
        print(f'  {sn}: {len(records)} PDAs')
        new_records.extend(records)

    all_records = existing_1 + new_records
    with open(OUTPUT_PATH, 'w', encoding='utf-8') as f:
        json.dump(all_records, f, ensure_ascii=False, indent=2)

    print(f'\nTotal: {len(all_records)} registros')
    print(f'  1\xb0 Secundaria: {len(existing_1)}')
    print(f'  2\xb0 Secundaria: {len(new_records)}')

if __name__ == '__main__':
    main()
